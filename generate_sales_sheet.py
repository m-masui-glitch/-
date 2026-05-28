#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date, timedelta

# ==============================
# 設定
# ==============================
VISIT_CYCLE = 10          # 平均来店周期（日）
ANALYSIS_MONTH_START = date(2026, 5, 1)
ANALYSIS_MONTH_END   = date(2026, 5, 31)

# ==============================
# 保有顧客データ（CSVより抽出）
# id, 氏名, 券種, 最終来店日, 残り回数, メモ
# ==============================
RAW_DATA = {
    "12回券": [
        (2014, "サイトウ　ヨシエ",      None,              2,  "都度"),
        (2035, "ハギワラ　マユミ",      None,              4,  "都度"),
        (2039, "タカナシ　エイイチ",    None,              0,  "都度"),
        (2043, "ナカタ　ヨウコ",        None,              0,  "都度"),
        (2348, "コダマ　ショウ",        date(2026,5,24),   4,  ""),
        (2450, "アリフク　ジュン",      date(2026,5,18),   7,  ""),
        (2917, "クス　フミエ",          date(2026,5,26),   8,  ""),
        (2937, "カワゾエ　イチロウ",    date(2026,5,24),   4,  ""),
        (2941, "エグチ　ヒロヒト",      date(2026,4,29),   0,  ""),
        (2982, "タカハシ　エミコ",      date(2026,5,27),   1,  ""),
        (3006, "ワダ　シュウセイ",      date(2026,5,22),  12,  ""),
        (3007, "イチカワ　タカヒロ",    date(2026,5,8),    2,  ""),
        (3061, "マスダ　ハツミ",        date(2026,5,22),   6,  ""),
        (3089, "デガワ　ハルト",        date(2026,5,19),   8,  ""),
    ],
    "8回券": [
        (5294, "クマガイ　ケン",        date(2026,5,15),   7,  ""),
        (2022, "オオハシ　タツエ",      None,              1,  "都度"),
        (2025, "アライ　シンスケ",      None,              0,  "都度"),
        (2034, "タナベ　ナナエ",        None,              3,  "都度"),
        (2037, "タナベ　ヒロユキ",      None,              1,  "都度"),
        (2063, "ハツタ　ユウスケ",      None,              1,  "都度"),
        (2065, "スズキ　ヒサキ",        None,              1,  "都度"),
        (2068, "エンドウ　ユキエ",      None,              0,  "都度"),
        (2069, "タナカ　ヨネコ",        None,              0,  "都度"),
        (2098, "エビサワ　トモコ",      date(2026,5,2),    5,  ""),
        (2148, "スヤマ　ヨシユキ",      date(2026,5,16),   8,  ""),
        (2228, "イトウ　トシアキ",      date(2026,5,18),   1,  ""),
        (2301, "キヨカワ　タエ",        date(2026,5,26),   5,  ""),
        (2308, "オオヤマ　シズオ",      date(2026,5,2),    8,  ""),
        (2365, "カサハラ　リョウコ",    date(2026,4,26),   1,  ""),
        (2377, "イシモリ　キヨタカ",    date(2026,5,16),   6,  ""),
        (2393, "タケイ　ミツヒロ",      date(2026,5,16),   4,  ""),
        (2459, "サイトウ　マサヒロ",    date(2026,5,24),   2,  ""),
        (2520, "ニッタ　ユウヤ",        date(2026,5,23),   6,  ""),
        (2537, "タニグチ　アキカズ",    date(2026,5,24),   4,  ""),
        (2597, "サイタ　アサヒ",        date(2026,3,20),   6,  ""),
        (2649, "イワモト　タカヒト",    date(2026,5,10),   4,  ""),
        (2761, "ノダ　カズエ",          date(2026,5,10),   6,  ""),
        (2797, "ムラマツ　ミホ",        date(2026,5,6),    2,  ""),
        (2809, "オカザキ　ミツル",      date(2026,5,19),   1,  "3回券2/3使用中"),
        (2864, "オヌキ　コウジ",        date(2026,5,17),   4,  ""),
        (2875, "ムカイ　ヒロト",        date(2026,5,23),   6,  ""),
        (2911, "ソガ　アキノリ",        date(2026,5,16),   1,  "3回券1/3使用中"),
        (3055, "ツマガリ　タイキ",      date(2026,5,10),   2,  ""),
        (3057, "コヤマ　ミキ",          date(2026,5,15),   3,  ""),
        (3068, "タカミ　トシユキ",      date(2026,4,11),   1,  ""),
        (3091, "クラタ　ミノリ",        date(2026,5,17),   6,  ""),
        (3096, "イワサキ　ミサコ",      date(2026,5,19),   6,  ""),
        (3098, "ヤマモト　ヨシノリ",    date(2026,5,26),   4,  ""),
        (3102, "タカツカ　ミホ",        date(2026,5,15),   8,  ""),
        (3103, "トミナガ　タイチ",      date(2026,5,23),   7,  ""),
        (3104, "タナカ　チサト",        date(2026,5,22),   7,  ""),
    ],
    "4回券": [
        (2007, "ツチダ　トモミ",        date(2026,5,17),   4,  ""),
        (5020, "ハヤシ　タツヤ",        date(2026,3,23),   2,  ""),
        (2021, "オオガヤ　シュンジ",    None,              0,  "都度"),
        (2026, "ミズキ　ヨシコ",        None,              3,  "都度"),
        (2027, "ヤナモト　カズキ",      None,              3,  "都度"),
        (2030, "カネコ　コウタロウ",    None,              0,  "都度"),
        (2040, "ハガ　レイコ",          None,              1,  "都度"),
        (2041, "カイ　カヨ",            None,              1,  "都度"),
        (2042, "サトウ　トオル",        None,              0,  "都度"),
        (2044, "ヤマモト　ユウスケ",    None,              0,  "都度"),
        (2046, "イシカワ　シュウヘイ",  None,              1,  "都度"),
        (2066, "テラジマ　ヨシアキ",    None,              3,  "都度"),
        (5142, "カワツ　チアキ",        date(2026,4,15),   0,  ""),
        (2291, "ササキ　クミ",          date(2026,5,17),   1,  ""),
        (2337, "フジイ　カヨコ",        date(2026,5,13),   2,  ""),
        (2368, "イマイ　マサヒロ",      date(2026,5,17),   1,  ""),
        (2394, "コイズミ　ヒロノリ",    date(2026,5,11),   3,  ""),
        (2396, "スガ　ヒロカズ",        date(2026,5,18),   0,  ""),
        (2453, "ハセガワ　ミツアキ",    date(2026,5,5),    3,  ""),
        (2526, "マツダ　リラ",          date(2026,4,27),   2,  ""),
        (2563, "トキトウ　フミト",      date(2026,3,21),   1,  ""),
        (2707, "サトウ　ヨシヒコ",      date(2026,5,8),    2,  ""),
        (2744, "ノグチ　ユウリ",        date(2026,5,4),    1,  ""),
        (2747, "イノウエ　エミ",        date(2026,5,20),   3,  ""),
        (2793, "クサジマ　アレン",      date(2026,4,6),    3,  ""),
        (2822, "アオヤマ　ワカコ",      date(2026,5,10),   3,  ""),
        (2830, "シラホ　ケンイチ",      date(2026,5,16),   4,  ""),
        (2844, "クワバラ　ユウジ",      date(2026,5,6),    2,  ""),
        (2853, "ミヤウチ　ヤスコ",      date(2026,5,15),   1,  ""),
        (2858, "イケダ　ヒロコ",        date(2026,4,27),   2,  ""),
        (2866, "コイケ　アンナ",        date(2026,4,27),   2,  ""),
        (2871, "エグチ　ツギコ",        date(2026,5,8),    2,  ""),
        (2929, "タケダ　アワカ",        date(2026,4,25),   3,  ""),
        (2961, "ヨコハタ　カオリ",      date(2026,5,6),    4,  ""),
        (2974, "セイノ　アキヒコ",      date(2026,5,15),   3,  ""),
        (2975, "ヤギ　カナコ",          date(2026,5,27),   0,  ""),
        (3010, "クワタ　キミコ",        date(2026,5,18),   1,  ""),
        (3025, "ツルタ　カズコ",        date(2026,5,3),    4,  ""),
        (3026, "クラ　ユウジ",          date(2026,5,22),   3,  ""),
        (3039, "ミヤケ　カヨコ",        date(2026,2,21),   1,  ""),
        (3048, "イノウエ　カズヒロ",    date(2026,4,25),   4,  ""),
        (3051, "オカノ　ハンナ",        date(2026,5,4),    1,  ""),
        (3065, "ヨシノ　マサヨシ",      date(2026,5,3),    2,  ""),
        (3072, "フジイ　ミク",          date(2026,2,16),   3,  ""),
        (3078, "オオジ　ミユキ",        date(2026,5,22),   1,  ""),
        (3095, "ババ　ナツホ",          date(2026,5,23),   3,  ""),
        (3099, "ヨシノ　ショウコ",      date(2026,5,8),    3,  ""),
    ],
    "3回券": [
        (2023, "ヒガ　ケイスケ",        None,              1,  "都度"),
        (2024, "ヤマゾエ　タツロウ",    None,              0,  "都度"),
        (2028, "ニシ　トモヒロ",        None,              0,  "都度"),
        (2029, "カネコ　ユミ",          None,              0,  "都度"),
        (2031, "ヒラマツ　シュウイチ",  None,              1,  "都度"),
        (2032, "センダ　マコト",        None,              0,  "都度"),
        (2038, "フナバシ　ミチコ",      None,              0,  "都度"),
        (2047, "ヒラバヤシ　ヨシミ",    None,              0,  "都度"),
        (2050, "タカナシ　チエコ",      None,              1,  "都度"),
        (2052, "ナカムラ　トシアキ",    None,              0,  "都度"),
        (2055, "カミヤマ　サトル",      None,              0,  "都度"),
        (2056, "ソエカワ　ゴロウ",      None,              0,  "都度"),
        (2059, "ノガワ　マユミ",        None,              1,  "都度"),
        (2060, "サトウ　リキ",          None,              0,  "都度"),
        (2061, "サトウ　ケイ",          None,              0,  "都度"),
        (2064, "ヤマグチ　ソノコ",      None,              0,  "都度"),
        (2071, "ゴトウ　ナツミ",        None,              0,  "都度"),
        (2073, "サカモト　カツミ",      None,              1,  "都度"),
        (2074, "サイ　トモヤ",          None,              2,  "都度"),
        (2075, "オハラ　トモヒロ",      None,              0,  "都度"),
        (2889, "モトハシ　ケンイチ",    None,              0,  "都度"),
        (3093, "ムネヒサ　ユリ",        date(2026,4,26),   1,  ""),
        (3094, "イシヤマ　サキ",        date(2026,5,10),   1,  ""),
        (3108, "ナカジマ　ユウコ",      date(2026,4,29),   2,  ""),
        (3111, "イトウ　ケンタ",        date(2026,5,24),   1,  ""),
        (3112, "ノグチ　トシキ",        date(2026,5,24),   1,  ""),
        (3113, "マツイ　サアヤ",        date(2026,5,23),   1,  ""),
        (3114, "ツチヤ　ナツコ",        date(2026,5,19),   2,  ""),
        (3115, "アラカワ　ダイスケ",    date(2026,5,8),    3,  ""),
        (3117, "ヤマモト　サオリ",      date(2026,5,27),   2,  ""),
        (3118, "トキダ　ジュンイチ",    date(2026,5,16),   3,  ""),
        (3120, "ヨネヤマ　コウイチ",    date(2026,5,26),   2,  ""),
        (3123, "オオヒラ　タクミ",      date(2026,5,26),   2,  ""),
    ],
}

# ==============================
# スタイル定義
# ==============================
def border_thin():
    side = Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

HEADER_FILL   = fill("1F4E79")
HEADER_FONT   = Font(name="游ゴシック", bold=True, color="FFFFFF", size=10)
SUB_FILL_12   = fill("BDD7EE")   # 12回券 ライトブルー
SUB_FILL_8    = fill("C6EFCE")   # 8回券  ライトグリーン
SUB_FILL_4    = fill("FFEB9C")   # 4回券  ライトイエロー
SUB_FILL_3    = fill("FCE4D6")   # 3回券  ライトオレンジ
OVERDUE_FILL  = fill("FF0000")   # 期限切れ 赤
OVERDUE_FONT  = Font(name="游ゴシック", bold=True, color="FFFFFF", size=10)
URGENT_FILL   = fill("FFC7CE")   # 要確認  薄赤
INPUT_FILL    = fill("FFFFC0")   # 入力セル 薄黄
CALC_FILL     = fill("F2F2F2")   # 計算セル グレー
SECTION_FILL  = fill("2E75B6")   # セクション見出し
SECTION_FONT  = Font(name="游ゴシック", bold=True, color="FFFFFF", size=11)
TOTAL_FILL    = fill("1F4E79")
TOTAL_FONT    = Font(name="游ゴシック", bold=True, color="FFFFFF", size=12)
BODY_FONT     = Font(name="游ゴシック", size=10)
BOLD_FONT     = Font(name="游ゴシック", bold=True, size=10)
CENTER        = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT          = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT         = Alignment(horizontal="right",  vertical="center")

# ==============================
# ユーティリティ
# ==============================
def calc_renewal_date(last_visit, remaining):
    if last_visit is None:
        return None
    if remaining == 0:
        return last_visit  # 既に0枚→即更新必要
    return last_visit + timedelta(days=remaining * VISIT_CYCLE)

def classify(rd, last_visit, remaining):
    if rd is None:
        return "算出不可（都度）"
    if remaining == 0 and last_visit is not None and last_visit < ANALYSIS_MONTH_START:
        return "要確認（既に残0）"
    if remaining == 0 and last_visit is not None:
        return "当月更新必要（残0）"
    if rd < ANALYSIS_MONTH_START:
        return "要確認（更新期限超過）"
    if ANALYSIS_MONTH_START <= rd <= ANALYSIS_MONTH_END:
        return "当月更新見込み"
    return "翌月以降"

TYPE_FILLS = {
    "12回券": SUB_FILL_12,
    "8回券":  SUB_FILL_8,
    "4回券":  SUB_FILL_4,
    "3回券":  SUB_FILL_3,
}

# ==============================
# 顧客リスト解析
# ==============================
all_processed = []
for ctype, rows in RAW_DATA.items():
    for (cid, name, last_visit, remaining, notes) in rows:
        rd = calc_renewal_date(last_visit, remaining)
        status = classify(rd, last_visit, remaining)
        all_processed.append({
            "id": cid, "name": name, "type": ctype,
            "last_visit": last_visit, "remaining": remaining,
            "renewal_date": rd, "status": status, "notes": notes,
        })

# 当月更新見込み（都度除く）
may_list = [c for c in all_processed
            if "当月" in c["status"]]

# 過去期限超過（要確認）
overdue_list = [c for c in all_processed
                if "要確認" in c["status"]]

# 翌月以降
future_list = [c for c in all_processed
               if c["status"] == "翌月以降"]

# カウント（当月更新見込みのみ）
may_counts = {}
for ctype in ["12回券","8回券","4回券","3回券"]:
    may_counts[ctype] = len([c for c in may_list if c["type"] == ctype])

# ==============================
# Excel ワークブック作成
# ==============================
wb = openpyxl.Workbook()

# ==================== Sheet 1: 当月更新見込み客リスト ====================
ws1 = wb.active
ws1.title = "①当月更新見込みリスト"

def set_cell(ws, row, col, value, font=None, fill_style=None, alignment=None, number_format=None, border=True):
    c = ws.cell(row=row, column=col, value=value)
    if font:        c.font = font
    if fill_style:  c.fill = fill_style
    if alignment:   c.alignment = alignment
    if number_format: c.number_format = number_format
    if border:      c.border = border_thin()
    return c

# タイトル
ws1.merge_cells("A1:H1")
c = ws1["A1"]
c.value = "2026年5月 回数券更新見込み客リスト（平均来店周期10日ベース）"
c.font = Font(name="游ゴシック", bold=True, size=14, color="1F4E79")
c.alignment = CENTER
ws1.row_dimensions[1].height = 30

ws1.merge_cells("A2:H2")
c = ws1["A2"]
c.value = f"抽出基準：更新予定日 = 最終来店日 + 残り回数 × {VISIT_CYCLE}日 ／ 対象月：2026年5月1日〜5月31日"
c.font = Font(name="游ゴシック", size=9, color="595959")
c.alignment = LEFT

# ヘッダー
headers = ["顧客番号","氏名","券種","最終来店日","残り回数","更新予定日","ステータス","備考"]
col_widths = [10, 18, 8, 13, 8, 13, 22, 20]
for i, (h, w) in enumerate(zip(headers, col_widths), 1):
    set_cell(ws1, 3, i, h, font=HEADER_FONT, fill_style=HEADER_FILL, alignment=CENTER)
    ws1.column_dimensions[get_column_letter(i)].width = w
ws1.row_dimensions[3].height = 22

row = 4
sections = [
    ("当月更新見込み（残0＋当月期限）", may_list),
    ("要確認：更新期限超過（4月以前）", overdue_list),
]

for section_title, customers in sections:
    # セクション見出し
    ws1.merge_cells(f"A{row}:H{row}")
    c = ws1[f"A{row}"]
    c.value = f"▼ {section_title}　（{len(customers)}名）"
    c.font = SECTION_FONT
    c.fill = SECTION_FILL
    c.alignment = LEFT
    c.border = border_thin()
    ws1.row_dimensions[row].height = 20
    row += 1

    for ctype in ["12回券","8回券","4回券","3回券"]:
        type_customers = sorted(
            [c for c in customers if c["type"] == ctype],
            key=lambda x: x["renewal_date"] if x["renewal_date"] else date.max
        )
        if not type_customers:
            continue

        # 券種サブヘッダー
        ws1.merge_cells(f"A{row}:H{row}")
        c = ws1[f"A{row}"]
        c.value = f"  【{ctype}】　{len(type_customers)}名"
        c.font = Font(name="游ゴシック", bold=True, size=10)
        c.fill = TYPE_FILLS[ctype]
        c.alignment = LEFT
        c.border = border_thin()
        ws1.row_dimensions[row].height = 18
        row += 1

        for cust in type_customers:
            st = cust["status"]
            row_fill = TYPE_FILLS[ctype]
            row_font = BODY_FONT
            if "超過" in st or "既に残0" in st:
                row_fill = OVERDUE_FILL
                row_font = OVERDUE_FONT
            elif "当月更新必要" in st:
                row_fill = URGENT_FILL
                row_font = Font(name="游ゴシック", bold=True, size=10)

            lv = cust["last_visit"].strftime("%Y/%m/%d") if cust["last_visit"] else "—"
            rv_str = cust["renewal_date"].strftime("%Y/%m/%d") if cust["renewal_date"] else "—"
            vals = [
                cust["id"], cust["name"], cust["type"],
                lv, cust["remaining"], rv_str, st, cust["notes"]
            ]
            for col_i, v in enumerate(vals, 1):
                align = CENTER if col_i in (1,3,4,5,6) else LEFT
                set_cell(ws1, row, col_i, v, font=row_font, fill_style=row_fill, alignment=align)
            ws1.row_dimensions[row].height = 18
            row += 1

    row += 1  # セクション間スペース

# ==================== Sheet 2: 売上目標設定シート ====================
ws2 = wb.create_sheet("②売上目標設定シート")

ws2.column_dimensions["A"].width = 28
ws2.column_dimensions["B"].width = 16
ws2.column_dimensions["C"].width = 14
ws2.column_dimensions["D"].width = 16
ws2.column_dimensions["E"].width = 16
ws2.column_dimensions["F"].width = 14

# ---- タイトル ----
ws2.merge_cells("A1:F1")
c = ws2["A1"]
c.value = "売上目標設定シート　／　2026年5月"
c.font = Font(name="游ゴシック", bold=True, size=16, color="1F4E79")
c.alignment = CENTER
ws2.row_dimensions[1].height = 36

ws2.merge_cells("A2:F2")
c = ws2["A2"]
c.value = "※ 黄色セルに数字を入力すると自動計算されます"
c.font = Font(name="游ゴシック", size=9, color="FF0000", italic=True)
c.alignment = LEFT
ws2.row_dimensions[2].height = 18

# ---- 列ヘッダー ----
col_headers = ["項目", "更新見込み件数\n（変更可）", "回数券単価（円）\n【要入力】", "小計（円）", "備考", ""]
col_header_fills = [HEADER_FILL]*6
for i, h in enumerate(col_headers[:5], 1):
    set_cell(ws2, 3, i, h, font=HEADER_FONT, fill_style=HEADER_FILL, alignment=CENTER)
ws2.row_dimensions[3].height = 36

# ---- セクション1: 既存顧客 回数券更新見込み ----
ws2.merge_cells("A4:F4")
c = ws2["A4"]
c.value = "■ 既存顧客　回数券更新見込み（当月）"
c.font = SECTION_FONT
c.fill = SECTION_FILL
c.alignment = LEFT
c.border = border_thin()
ws2.row_dimensions[4].height = 22

coupon_rows = [
    ("12回券 更新見込み", may_counts["12回券"], SUB_FILL_12, 5),
    ("8回券 更新見込み",  may_counts["8回券"],  SUB_FILL_8,  6),
    ("4回券 更新見込み",  may_counts["4回券"],  SUB_FILL_4,  7),
    ("3回券 更新見込み",  may_counts["3回券"],  SUB_FILL_3,  8),
]

for label, cnt, row_fill, r in coupon_rows:
    # 項目名
    set_cell(ws2, r, 1, label, font=BOLD_FONT, fill_style=row_fill, alignment=LEFT)
    # 件数（変更可・黄色）
    c_cnt = ws2.cell(row=r, column=2, value=cnt)
    c_cnt.font = BOLD_FONT
    c_cnt.fill = INPUT_FILL
    c_cnt.alignment = CENTER
    c_cnt.border = border_thin()
    c_cnt.number_format = "0"
    # 単価（入力必須・黄色）
    c_price = ws2.cell(row=r, column=3, value=0)
    c_price.font = Font(name="游ゴシック", bold=True, size=10, color="C00000")
    c_price.fill = INPUT_FILL
    c_price.alignment = CENTER
    c_price.border = border_thin()
    c_price.number_format = "#,##0"
    # 小計（自動計算）
    col_b = get_column_letter(2)
    col_c = get_column_letter(3)
    c_sub = ws2.cell(row=r, column=4)
    c_sub.value = f"={col_b}{r}*{col_c}{r}"
    c_sub.font = BOLD_FONT
    c_sub.fill = CALC_FILL
    c_sub.alignment = RIGHT
    c_sub.border = border_thin()
    c_sub.number_format = "#,##0"
    # 備考
    memo = {"12回券 更新見込み":"12枚綴り","8回券 更新見込み":"8枚綴り",
            "4回券 更新見込み":"4枚綴り","3回券 更新見込み":"3枚綴り"}.get(label,"")
    set_cell(ws2, r, 5, memo, font=BODY_FONT, alignment=LEFT)
    ws2.row_dimensions[r].height = 22

# 既存小計
ws2.merge_cells("A9:C9")
c = ws2["A9"]
c.value = "既存顧客　小計"
c.font = Font(name="游ゴシック", bold=True, size=11, color="1F4E79")
c.fill = fill("DEEAF1")
c.alignment = RIGHT
c.border = border_thin()

c_sub9 = ws2.cell(row=9, column=4)
c_sub9.value = "=D5+D6+D7+D8"
c_sub9.font = Font(name="游ゴシック", bold=True, size=11, color="1F4E79")
c_sub9.fill = fill("DEEAF1")
c_sub9.alignment = RIGHT
c_sub9.border = border_thin()
c_sub9.number_format = "#,##0"

ws2.cell(row=9, column=5).border = border_thin()
ws2.row_dimensions[9].height = 24

# ---- セクション2: 新規顧客 ----
ws2.merge_cells("A10:F10")
c = ws2["A10"]
c.value = "■ 新規顧客"
c.font = SECTION_FONT
c.fill = SECTION_FILL
c.alignment = LEFT
c.border = border_thin()
ws2.row_dimensions[10].height = 22

# 新規件数
set_cell(ws2, 11, 1, "新規顧客数", font=BOLD_FONT, fill_style=fill("FFF2CC"), alignment=LEFT)
c_new = ws2.cell(row=11, column=2, value=0)
c_new.font = Font(name="游ゴシック", bold=True, size=10, color="C00000")
c_new.fill = INPUT_FILL
c_new.alignment = CENTER
c_new.border = border_thin()
c_new.number_format = "0"

c_price_new = ws2.cell(row=11, column=3, value=0)
c_price_new.font = Font(name="游ゴシック", bold=True, size=10, color="C00000")
c_price_new.fill = INPUT_FILL
c_price_new.alignment = CENTER
c_price_new.border = border_thin()
c_price_new.number_format = "#,##0"

set_cell(ws2, 11, 4, None, font=BODY_FONT, fill_style=CALC_FILL, alignment=RIGHT)
ws2.cell(row=11, column=4).value = "=B11*C11"
ws2.cell(row=11, column=4).number_format = "#,##0"
ws2.cell(row=11, column=4).font = BOLD_FONT
ws2.cell(row=11, column=4).fill = CALC_FILL
ws2.cell(row=11, column=4).border = border_thin()

set_cell(ws2, 11, 5, "初回チケット平均単価", font=BODY_FONT, alignment=LEFT)
ws2.row_dimensions[11].height = 22

# 新規小計
ws2.merge_cells("A12:C12")
c = ws2["A12"]
c.value = "新規顧客　小計"
c.font = Font(name="游ゴシック", bold=True, size=11, color="1F4E79")
c.fill = fill("E2EFDA")
c.alignment = RIGHT
c.border = border_thin()

c_sub12 = ws2.cell(row=12, column=4)
c_sub12.value = "=D11"
c_sub12.font = Font(name="游ゴシック", bold=True, size=11, color="1F4E79")
c_sub12.fill = fill("E2EFDA")
c_sub12.alignment = RIGHT
c_sub12.border = border_thin()
c_sub12.number_format = "#,##0"
ws2.cell(row=12, column=5).border = border_thin()
ws2.row_dimensions[12].height = 24

# ---- 合計 ----
ws2.row_dimensions[13].height = 8  # スペース

ws2.merge_cells("A14:C14")
c = ws2["A14"]
c.value = "▶ 当月　売上目標（最低見込み）"
c.font = TOTAL_FONT
c.fill = TOTAL_FILL
c.alignment = RIGHT
c.border = border_thin()

c_total = ws2.cell(row=14, column=4)
c_total.value = "=D9+D12"
c_total.font = Font(name="游ゴシック", bold=True, size=14, color="FFFFFF")
c_total.fill = TOTAL_FILL
c_total.alignment = RIGHT
c_total.border = border_thin()
c_total.number_format = "#,##0"
ws2.cell(row=14, column=5).border = border_thin()
ws2.row_dimensions[14].height = 32

# 注記
ws2.merge_cells("A15:F15")
c = ws2["A15"]
c.value = "※ この数値は「確実に見込める最低限の売上」です。都度来院顧客・新規のアップセルなど上振れ分は含まれません。"
c.font = Font(name="游ゴシック", size=9, color="595959", italic=True)
c.alignment = LEFT
ws2.row_dimensions[15].height = 18

# ---- サマリーテーブル ----
ws2.row_dimensions[17].height = 10
ws2.merge_cells("A17:F17")
c = ws2["A17"]
c.value = "■ 当月更新見込み　件数サマリー（参考）"
c.font = Font(name="游ゴシック", bold=True, size=12, color="1F4E79")
c.alignment = LEFT
ws2.row_dimensions[17].height = 22

summary_headers = ["券種", "当月更新見込み", "要確認（期限超過）", "翌月以降", "都度（算出不可）", "合計"]
for i, h in enumerate(summary_headers, 1):
    set_cell(ws2, 18, i, h, font=HEADER_FONT, fill_style=HEADER_FILL, alignment=CENTER)
ws2.row_dimensions[18].height = 22

for r_i, ctype in enumerate(["12回券","8回券","4回券","3回券"], 19):
    may_c   = len([c for c in all_processed if c["type"]==ctype and "当月" in c["status"]])
    over_c  = len([c for c in all_processed if c["type"]==ctype and "要確認" in c["status"]])
    fut_c   = len([c for c in all_processed if c["type"]==ctype and c["status"]=="翌月以降"])
    tsudo_c = len([c for c in all_processed if c["type"]==ctype and "算出不可" in c["status"]])
    total_c = may_c + over_c + fut_c + tsudo_c
    rf = TYPE_FILLS[ctype]
    for col_i, v in enumerate([ctype, may_c, over_c, fut_c, tsudo_c, total_c], 1):
        set_cell(ws2, r_i, col_i, v, font=BOLD_FONT if col_i in (1,2,6) else BODY_FONT,
                 fill_style=rf if col_i==1 else (URGENT_FILL if col_i==2 else
                 (OVERDUE_FILL if col_i==3 else (CALC_FILL if col_i==6 else None))),
                 alignment=CENTER)
    ws2.row_dimensions[r_i].height = 20

# 合計行
totals = [
    len([c for c in all_processed if "当月" in c["status"]]),
    len([c for c in all_processed if "要確認" in c["status"]]),
    len([c for c in all_processed if c["status"]=="翌月以降"]),
    len([c for c in all_processed if "算出不可" in c["status"]]),
]
totals.append(sum(totals))
set_cell(ws2, 23, 1, "合計", font=BOLD_FONT, fill_style=fill("DEEAF1"), alignment=CENTER)
for col_i, v in enumerate(totals, 2):
    set_cell(ws2, 23, col_i, v, font=BOLD_FONT, fill_style=fill("DEEAF1"), alignment=CENTER)
ws2.row_dimensions[23].height = 22

# ---- 入力ガイド ----
ws2.row_dimensions[25].height = 8
ws2.merge_cells("A25:F25")
c = ws2["A25"]
c.value = "■ 入力手順"
c.font = Font(name="游ゴシック", bold=True, size=12, color="1F4E79")
c.alignment = LEFT
ws2.row_dimensions[25].height = 22

guides = [
    "① 各回数券の「単価」（C5〜C8）を入力　→ 既存顧客の小計が自動計算されます",
    "② 「新規顧客数」（B11）を実績または予測数に変更",
    "③ 「初回チケット平均単価」（C11）を入力",
    "④ B5〜B8 の「更新見込み件数」は変更可能（実態に合わせて上書きしてください）",
    "⑤ 売上目標セル（D14）に当月の最低見込み売上が表示されます",
]
for g_i, g in enumerate(guides, 26):
    ws2.merge_cells(f"A{g_i}:F{g_i}")
    c = ws2[f"A{g_i}"]
    c.value = g
    c.font = Font(name="游ゴシック", size=10)
    c.alignment = LEFT
    ws2.row_dimensions[g_i].height = 18

# ==================== Sheet 3: 全顧客ステータス ====================
ws3 = wb.create_sheet("③全顧客ステータス一覧")

ws3.column_dimensions["A"].width = 10
ws3.column_dimensions["B"].width = 18
ws3.column_dimensions["C"].width = 8
ws3.column_dimensions["D"].width = 13
ws3.column_dimensions["E"].width = 8
ws3.column_dimensions["F"].width = 13
ws3.column_dimensions["G"].width = 22
ws3.column_dimensions["H"].width = 18

ws3.merge_cells("A1:H1")
c = ws3["A1"]
c.value = "保有顧客　全ステータス一覧"
c.font = Font(name="游ゴシック", bold=True, size=14, color="1F4E79")
c.alignment = CENTER
ws3.row_dimensions[1].height = 28

headers3 = ["顧客番号","氏名","券種","最終来店日","残り回数","更新予定日","ステータス","備考"]
for i, h in enumerate(headers3, 1):
    set_cell(ws3, 2, i, h, font=HEADER_FONT, fill_style=HEADER_FILL, alignment=CENTER)
ws3.row_dimensions[2].height = 22

sorted_all = sorted(all_processed,
    key=lambda x: (
        0 if "要確認" in x["status"] else
        1 if "当月" in x["status"] else
        2 if x["status"] == "翌月以降" else 3,
        x["renewal_date"] if x["renewal_date"] else date.max,
        x["type"], x["id"]
    )
)

for r_i, cust in enumerate(sorted_all, 3):
    st = cust["status"]
    if "要確認" in st:
        rf, rft = OVERDUE_FILL, OVERDUE_FONT
    elif "当月更新必要" in st:
        rf, rft = URGENT_FILL, Font(name="游ゴシック", bold=True, size=10)
    elif "当月更新見込み" in st:
        rf, rft = TYPE_FILLS[cust["type"]], BOLD_FONT
    else:
        rf, rft = None, BODY_FONT

    lv = cust["last_visit"].strftime("%Y/%m/%d") if cust["last_visit"] else "—"
    rv = cust["renewal_date"].strftime("%Y/%m/%d") if cust["renewal_date"] else "—"
    vals = [cust["id"], cust["name"], cust["type"], lv, cust["remaining"], rv, st, cust["notes"]]
    for col_i, v in enumerate(vals, 1):
        align = CENTER if col_i in (1,3,4,5,6) else LEFT
        set_cell(ws3, r_i, col_i, v, font=rft, fill_style=rf, alignment=align)
    ws3.row_dimensions[r_i].height = 16

# ==============================
# 保存
# ==============================
output_path = "/home/user/-/売上目標設定シート_2026年5月.xlsx"
wb.save(output_path)
print(f"✅ 保存完了: {output_path}")

# サマリー表示
print("\n=== 当月更新見込み（2026年5月） ===")
for ctype in ["12回券","8回券","4回券","3回券"]:
    cnt = may_counts[ctype]
    print(f"  {ctype}: {cnt}名")
print(f"\n  合計: {sum(may_counts.values())}名")
print("\n=== 要確認（更新期限超過） ===")
for c in overdue_list:
    lv = c['last_visit'].strftime('%m/%d') if c['last_visit'] else '—'
    rv = c['renewal_date'].strftime('%m/%d') if c['renewal_date'] else '—'
    print(f"  [{c['type']}] {c['name']} (最終:{lv} 残{c['remaining']}枚 → 更新予定:{rv}) {c['status']}")
